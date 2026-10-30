from __future__ import annotations

from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "tmp" / "resume-build"
BLACK = RGBColor(18, 24, 32)
MUTED = RGBColor(58, 67, 78)
ACCENT = RGBColor(15, 72, 94)


def set_cell_free_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    relation_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0F485E")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend([color, underline])
    run.append(properties)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def configure(document: Document, language: str) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.68)
    section.right_margin = Inches(0.68)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.2)
    normal.font.color.rgb = MUTED
    normal.paragraph_format.space_after = Pt(2.5)
    normal.paragraph_format.line_spacing = 1.03

    for style_name, size in [("Title", 20), ("Heading 1", 10.8), ("Heading 2", 10.2)]:
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = BLACK
        style.font.bold = True
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(6 if style_name == "Heading 1" else 2)
        style.paragraph_format.space_after = Pt(2)
    styles["Heading 1"].font.all_caps = True

    document.core_properties.title = (
        "Sebastián Astroza — General Software Engineer Resume"
        if language == "en"
        else "Sebastián Astroza — Currículum General de Ingeniería de Software"
    )
    document.core_properties.author = "Sebastián Andrés Astroza Aguayo"
    document.core_properties.subject = "Software engineering, systems, data and applied AI"
    document.core_properties.keywords = "software engineering, full-stack, Linux, data, automation, machine learning"
    document.core_properties.comments = "Professional portfolio resume — 2026"


def add_name_and_contact(document: Document, language: str) -> None:
    name = document.add_paragraph(style="Title")
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.paragraph_format.space_after = Pt(0)
    run = name.add_run("SEBASTIÁN ANDRÉS ASTROZA AGUAYO")
    run.font.color.rgb = BLACK

    role = document.add_paragraph()
    role.alignment = WD_ALIGN_PARAGRAPH.CENTER
    role.paragraph_format.space_after = Pt(2)
    role_run = role.add_run("Software Engineer | Full-Stack, Systems, Data & Applied AI" if language == "en" else "Ingeniero de Software | Full-Stack · Sistemas · Datos · IA")
    role_run.bold = True
    role_run.font.size = Pt(10.2)
    role_run.font.color.rgb = ACCENT

    contact = document.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_after = Pt(0)
    contact.add_run("Santiago, Chile  |  ")
    set_cell_free_hyperlink(contact, "sebastroza10@gmail.com", "mailto:sebastroza10@gmail.com")
    contact.add_run("  |  ")
    set_cell_free_hyperlink(contact, "linkedin.com/in/sebastian-a-677745276", "https://www.linkedin.com/in/sebastian-a-677745276/")

    links = document.add_paragraph()
    links.alignment = WD_ALIGN_PARAGRAPH.CENTER
    links.paragraph_format.space_after = Pt(4)
    set_cell_free_hyperlink(links, "github.com/sebastroza1", "https://github.com/sebastroza1")
    links.add_run("  |  ")
    set_cell_free_hyperlink(links, "sebastroza1.github.io", "https://sebastroza1.github.io")


def heading(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(text, style="Heading 1")
    paragraph.paragraph_format.keep_with_next = True
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), "B7C4CC")
    border.append(bottom)
    paragraph._p.get_or_add_pPr().append(border)


def add_role(document: Document, company: str, role: str, date: str, bullets: Iterable[str]) -> None:
    title = document.add_paragraph()
    title.paragraph_format.keep_with_next = True
    title.paragraph_format.space_after = Pt(0)
    run = title.add_run(f"{company} — {role}")
    run.bold = True
    run.font.color.rgb = BLACK
    run = title.add_run(f"  |  {date}")
    run.bold = True
    run.font.color.rgb = ACCENT
    for bullet in bullets:
        item = document.add_paragraph(style="List Bullet")
        item.paragraph_format.left_indent = Inches(0.18)
        item.paragraph_format.first_line_indent = Inches(-0.12)
        item.paragraph_format.space_after = Pt(1)
        item.add_run(bullet)


def add_project(document: Document, title: str, tech: str, description: str, url: str | None = None) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(1.2)
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run(f"{title} | {tech}: ")
    run.bold = True
    run.font.color.rgb = BLACK
    paragraph.add_run(description)
    if url:
        paragraph.add_run(" ")
        set_cell_free_hyperlink(paragraph, "Repository", url)


def add_skills(document: Document, rows: list[tuple[str, str]]) -> None:
    for label, values in rows:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(1)
        run = paragraph.add_run(f"{label}: ")
        run.bold = True
        run.font.color.rgb = BLACK
        paragraph.add_run(values)


def build_english() -> Document:
    doc = Document()
    configure(doc, "en")
    add_name_and_contact(doc, "en")
    heading(doc, "PROFESSIONAL SUMMARY")
    doc.add_paragraph(
        "Early-career software engineer with professional experience across full-stack delivery, Linux infrastructure, QA, automation and data integrations, complemented by independently validated engineering labs in applied AI, distributed systems and observability. Builds inspectable, documented software with explicit evidence boundaries; open to international and relocation opportunities."
    )

    heading(doc, "PROFESSIONAL EXPERIENCE")
    add_role(doc, "AutismoConecta", "Software Engineer — Cross-functional IT", "Sep 2026–Present", [
        "Contribute across full-stack development, Linux infrastructure, databases, automation, security, external integrations, administrative systems, QA, deployment validation and incident diagnosis in a Chilean social-impact technology context.",
        "Documented owner-provided local p75 evidence after web optimization: DOMContentLoaded −42.7%, First Contentful Paint −40.3% and initial bundle −34.4%; measurements are local, not real-user field data, and do not establish individual commit authorship.",
    ])
    add_role(doc, "TRC Recruitme", "QA, Automation & Data Integrations", "Aug 2025–Mar 2026", [
        "Supported technical, professional and university placement periods followed by continuity under contract; individual placement dates are not claimed.",
        "Worked on data extraction and transformation, Dropbox CV migration, import templates, n8n workflows, webhooks, chatbots, Selenium, Postman, Scrum and technical documentation.",
        "Owner-reported delivery: completed a migration task estimated at one month in approximately one week."
    ])

    heading(doc, "TECHNICAL SKILLS")
    add_skills(doc, [
        ("Software", "Python, PHP, JavaScript/TypeScript, Laravel, Angular, Next.js, Vue 3, FastAPI, REST APIs, C#, Java, Flutter"),
        ("Systems & DevOps", "Linux, Nginx, PHP-FPM, Docker, Git, GitHub Actions, CI/CD, Bash, PowerShell, networking, observability; Kubernetes coursework and labs"),
        ("Data", "PostgreSQL/PostGIS, MySQL, SQL, Firebase/Firestore, ETL, geospatial data, Pandas, NumPy"),
        ("Quality & AI", "pytest, Selenium, Postman, Playwright, web performance, PyTorch, TensorFlow/Keras, scikit-learn, MLflow, OpenCV"),
    ])

    heading(doc, "SELECTED PROJECTS")
    add_project(doc, "LandFinder Chile", "Vue 3, FastAPI, PostGIS, Leaflet", "Private work in progress implementing geospatial search, comparison, provenance-aware public-data connectors and deterministic territorial scoring; not deployed and not presented as a completed MVP.")
    add_project(doc, "Collaboration Operations Lab", "Python, Microsoft Graph/Entra concepts", "Public reconciliation lab with dry-run defaults, idempotent planning, least-privilege documentation and tested failure paths.", "https://github.com/sebastroza1/collaboration-ops-lab")
    add_project(doc, "ML Model Serving & Operations", "FastAPI, Docker, Prometheus", "Reproducible inference API with validation, health/metrics endpoints, tests and operational runbooks.", "https://github.com/sebastroza1/ml-model-serving-ops")
    add_project(doc, "Scientific Anomaly Detection", "PyTorch, scikit-learn", "Synthetic scientific-signal pipeline with deterministic generation, evaluation and clearly stated non-production limits.", "https://github.com/sebastroza1/scientific-anomaly-detection")
    add_project(doc, "C++ File Integrity Toolkit", "C++17, CMake", "Portable hashing and integrity-verification workflow with tests and CI evidence.", "https://github.com/sebastroza1/cpp-file-integrity-toolkit")

    heading(doc, "EDUCATION")
    doc.add_paragraph("Universidad Central de Chile — Civil Engineering in Computing and Informatics | Academic requirements completed 2026; formal degree conferral pending institutional confirmation. Teaching-assistant experience in Physics and Probability.")
    doc.add_paragraph("NeuroFAST | Selected for the first national stage of UDD Impacto Emprendedor 2026 and participated in formative activities; no competition win or clinical validation is claimed.")

    heading(doc, "CERTIFICATIONS")
    certs = doc.add_paragraph()
    certs.add_run("EF SET English Certificate — Global C1, Oct 2026 (credential: ").bold = True
    set_cell_free_hyperlink(certs, "cert.efset.org/m27XNx", "https://cert.efset.org/m27XNx")
    certs.add_run(")  |  LFS158: Introduction to Kubernetes — The Linux Foundation, Oct 2026 (course completion, owner-provided)")

    heading(doc, "LANGUAGES")
    doc.add_paragraph("Spanish — Native | English — Global C1 | French — Beginner | Portuguese — Basic")
    return doc


def build_spanish() -> Document:
    doc = Document()
    configure(doc, "es")
    add_name_and_contact(doc, "es")
    heading(doc, "PERFIL PROFESIONAL")
    doc.add_paragraph(
        "Ingeniero de software en etapa inicial con experiencia profesional en desarrollo full-stack, infraestructura Linux, QA, automatización e integraciones de datos, complementada por laboratorios de ingeniería validados de forma independiente en IA aplicada, sistemas distribuidos y observabilidad. Construye software inspeccionable y documentado, con límites explícitos de evidencia; disponible para oportunidades internacionales y reubicación."
    )

    heading(doc, "EXPERIENCIA PROFESIONAL")
    add_role(doc, "AutismoConecta", "Ingeniero de Software — TI transversal", "Sep 2026–Actualidad", [
        "Contribución en desarrollo full-stack, infraestructura Linux, bases de datos, automatización, seguridad, integraciones externas, sistemas administrativos, QA, validación de despliegues y diagnóstico de incidentes en un contexto tecnológico chileno de impacto social.",
        "Evidencia p75 local informada por el propietario tras optimización web: DOMContentLoaded −42,7 %, First Contentful Paint −40,3 % y bundle inicial −34,4 %; son mediciones locales, no datos de usuarios reales, y no establecen autoría individual de commits.",
    ])
    add_role(doc, "TRC Recruitme", "QA, Automatización e Integraciones de Datos", "Ago 2025–Mar 2026", [
        "Participación en prácticas técnica, profesional y universitaria, seguida por continuidad bajo contrato; no se afirman fechas individuales de cada práctica.",
        "Trabajo en extracción y transformación de datos, migración de CV desde Dropbox, plantillas de importación, n8n, webhooks, chatbots, Selenium, Postman, Scrum y documentación técnica.",
        "Entrega informada por el propietario: tarea de migración estimada en un mes completada en aproximadamente una semana."
    ])

    heading(doc, "HABILIDADES TÉCNICAS")
    add_skills(doc, [
        ("Software", "Python, PHP, JavaScript/TypeScript, Laravel, Angular, Next.js, Vue 3, FastAPI, APIs REST, C#, Java, Flutter"),
        ("Sistemas y DevOps", "Linux, Nginx, PHP-FPM, Docker, Git, GitHub Actions, CI/CD, Bash, PowerShell, redes, observabilidad; formación y laboratorios en Kubernetes"),
        ("Datos", "PostgreSQL/PostGIS, MySQL, SQL, Firebase/Firestore, ETL, datos geoespaciales, Pandas, NumPy"),
        ("Calidad e IA", "pytest, Selenium, Postman, Playwright, rendimiento web, PyTorch, TensorFlow/Keras, scikit-learn, MLflow, OpenCV"),
    ])

    doc.add_page_break()
    heading(doc, "PROYECTOS SELECCIONADOS")
    add_project(doc, "LandFinder Chile", "Vue 3, FastAPI, PostGIS, Leaflet", "Trabajo privado en desarrollo con búsqueda geoespacial, comparación, conectores de datos públicos con procedencia y scoring territorial determinista; no está desplegado ni se presenta como MVP completo.")
    add_project(doc, "Laboratorio de Operaciones Colaborativas", "Python, conceptos Microsoft Graph/Entra", "Laboratorio público de reconciliación con dry-run por defecto, planificación idempotente, mínimo privilegio y rutas de fallo probadas.", "https://github.com/sebastroza1/collaboration-ops-lab")
    add_project(doc, "Servicio y Operaciones de Modelos ML", "FastAPI, Docker, Prometheus", "API de inferencia reproducible con validación, endpoints de salud/métricas, pruebas y runbooks operativos.", "https://github.com/sebastroza1/ml-model-serving-ops")
    add_project(doc, "Detección de Anomalías Científicas", "PyTorch, scikit-learn", "Pipeline de señales científicas sintéticas con generación determinista, evaluación y límites no productivos explícitos.", "https://github.com/sebastroza1/scientific-anomaly-detection")
    add_project(doc, "Toolkit de Integridad de Archivos C++", "C++17, CMake", "Flujo portable de hashing y verificación de integridad con pruebas y evidencia CI.", "https://github.com/sebastroza1/cpp-file-integrity-toolkit")

    heading(doc, "EDUCACIÓN")
    doc.add_paragraph("Universidad Central de Chile — Ingeniería Civil en Computación e Informática | Requisitos académicos completados en 2026; otorgamiento formal del título pendiente de confirmación institucional. Experiencia como ayudante en Física y Probabilidad.")
    doc.add_paragraph("NeuroFAST | Seleccionado en la primera etapa nacional de Impacto Emprendedor UDD 2026 y participante en actividades formativas; no se afirma premio ni validación clínica.")

    heading(doc, "CERTIFICACIONES")
    certs = doc.add_paragraph()
    certs.add_run("EF SET English Certificate — C1 global, oct 2026 (credencial: ").bold = True
    set_cell_free_hyperlink(certs, "cert.efset.org/m27XNx", "https://cert.efset.org/m27XNx")
    certs.add_run(")  |  LFS158: Introduction to Kubernetes — The Linux Foundation, oct 2026 (curso completado, información del propietario)")

    heading(doc, "IDIOMAS")
    doc.add_paragraph("Español — Nativo | Inglés — C1 global | Francés — Principiante | Portugués — Básico")
    return doc


def main() -> None:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    files = {
        "Sebastian_Astroza_Software_Engineer_General_2026.docx": build_english(),
        "Sebastian_Astroza_CV_General_2026_ES.docx": build_spanish(),
    }
    for filename, document in files.items():
        path = OUTPUT / filename
        document.save(path)
        print(path)


if __name__ == "__main__":
    main()
